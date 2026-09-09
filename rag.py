import base64
import hashlib
import time
from pathlib import Path
from openai import OpenAI
from config import MODEL_EMBEDDING, MODEL_OCR

import chromadb
from chromadb.utils import embedding_functions

try:
    import pymupdf as fitz  # PyMuPDF
    HAS_PDF = True
except ImportError:
    try:
        import fitz  # PyMuPDF (เวอร์ชันเก่า)
        HAS_PDF = True
    except ImportError:
        HAS_PDF = False

try:
    from docx import Document as DocxDocument
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    from pptx import Presentation
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False


OCR_PROMPT = ("อ่านและถอดข้อความทั้งหมดในภาพนี้ออกมาตามลำดับที่ปรากฏ "
              "ห้ามสรุปหรือแปล ให้ตอบเฉพาะข้อความที่อ่านได้เท่านั้น")


def _ocr_page(client: OpenAI, img_b64: str, retries: int = 4):
    """เรียก OCR 1 หน้า พร้อม retry แบบ backoff (เผื่อโดน rate-limit 429/5xx)"""
    delay = 5
    for attempt in range(retries + 1):
        try:
            return client.chat.completions.create(
                model=MODEL_OCR,
                messages=[{
                    "role": "user",
                    "content": [
                        {"type": "text", "text": OCR_PROMPT},
                        {"type": "image_url",
                         "image_url": {"url": f"data:image/png;base64,{img_b64}"}}
                    ]
                }]
            )
        except Exception as e:
            msg = str(e)
            transient = any(c in msg for c in ("429", "500", "502", "503", "timeout", "Provider returned error"))
            if attempt >= retries or not transient:
                raise
            print(f"      รอ {delay}s แล้วลองใหม่ (ครั้งที่ {attempt + 1}): {msg[:80]}")
            time.sleep(delay)
            delay = min(delay * 2, 60)


def ocr_pdf(file_path: str, client: OpenAI, dpi: int = 200, cost_tracker=None) -> str:
    """แปลงแต่ละหน้า PDF เป็นรูปภาพแล้วให้ LLM อ่านข้อความ (OCR)"""
    if not HAS_PDF:
        raise ImportError("ติดตั้ง pymupdf ก่อน: pip install pymupdf")

    doc   = fitz.open(file_path)
    pages = []
    n_pages       = len(doc)
    in_tok, out_tok = 0, 0

    for i, page in enumerate(doc):
        pix       = page.get_pixmap(dpi=dpi)
        img_b64   = base64.b64encode(pix.tobytes("png")).decode("utf-8")

        response = _ocr_page(client, img_b64)
        pages.append(response.choices[0].message.content or "")
        print(f"    หน้า {i + 1}/{n_pages}")

        usage = getattr(response, "usage", None)
        if usage:
            in_tok  += getattr(usage, "prompt_tokens", 0) or 0
            out_tok += getattr(usage, "completion_tokens", 0) or 0

    doc.close()

    if cost_tracker is not None and (in_tok or out_tok):
        cost_tracker.track_ocr(in_tok, out_tok)

    return "\n\n".join(pages)


def pdf_to_txt(lesson_path, client: OpenAI, dpi: int = 200, cost_tracker=None) -> list[Path]:
    """OCR ไฟล์ PDF ทุกไฟล์ในโฟลเดอร์บทเรียนให้กลายเป็นไฟล์ .txt ข้างๆ กัน
    ข้ามไฟล์ที่มี .txt (ไม่ว่าง) อยู่แล้ว คืนค่า path ของไฟล์ .txt ทั้งหมด"""
    lesson_path = Path(lesson_path)
    pdfs = sorted(
        f for f in lesson_path.iterdir()
        if f.is_file() and f.suffix.lower() == ".pdf"
    )

    if not pdfs:
        print("  ไม่พบไฟล์ PDF ในโฟลเดอร์นี้")
        return []

    if not HAS_PDF:
        raise ImportError("ติดตั้ง pymupdf ก่อน: pip install pymupdf")

    txt_files = []
    for pdf in pdfs:
        txt_path = pdf.with_suffix(".txt")

        if txt_path.exists() and txt_path.stat().st_size > 0:
            print(f"  ข้าม {pdf.name} (มี {txt_path.name} อยู่แล้ว)")
            txt_files.append(txt_path)
            continue

        print(f"  OCR: {pdf.name} → {txt_path.name}")
        text = ocr_pdf(str(pdf), client, dpi=dpi, cost_tracker=cost_tracker)
        txt_path.write_text(text, encoding="utf-8")
        print(f"    เขียน {len(text)} ตัวอักษร")

        if not text.strip():
            print(f"    ⚠️  OCR ไม่ได้ข้อความจาก {pdf.name}")

        txt_files.append(txt_path)

    return txt_files


def extract_text(file_path: str, client: OpenAI = None, cost_tracker=None) -> str:
    path = Path(file_path)
    ext  = path.suffix.lower()

    if ext in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="ignore")

    elif ext == ".pdf":
        if client is None:
            raise ValueError("ต้องส่ง client (OpenAI/OpenRouter) เพื่อทำ OCR ไฟล์ PDF ด้วย Qwen")
        return ocr_pdf(file_path, client, cost_tracker=cost_tracker)

    elif ext == ".docx":
        if not HAS_DOCX:
            raise ImportError("ติดตั้ง python-docx ก่อน: pip install python-docx")
        doc = DocxDocument(file_path)
        return "\n".join(p.text for p in doc.paragraphs)

    elif ext == ".pptx":
        if not HAS_PPTX:
            raise ImportError("ติดตั้ง python-pptx ก่อน: pip install python-pptx")
        prs   = Presentation(file_path)
        texts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    texts.append(shape.text)
        return "\n".join(texts)

    else:
        raise ValueError(f"ไม่รองรับไฟล์ประเภท: {ext}")


def chunk_text(text: str, chunk_size: int = 1000, overlap: int = 120) -> list[str]:
    """แบ่งข้อความเป็น chunk ตามจำนวนตัวอักษร

    ใช้ตัวอักษรแทนคำ เพราะภาษาไทยไม่มีช่องว่างระหว่างคำ การใช้ text.split()
    จะได้ "คำ" ที่ยาวมากจน chunk เกิน token limit ของ embedding model
    """
    text = text.strip()
    if not text:
        return []

    chunks = []
    n      = len(text)
    start  = 0

    while start < n:
        end = min(start + chunk_size, n)

        if end < n:
            # ถอยไปหาจุดขึ้นบรรทัด/ช่องว่างที่ใกล้ที่สุด เพื่อไม่ตัดกลางประโยค
            window = text[start:end]
            cut    = max(window.rfind("\n"), window.rfind(" "))
            if cut > chunk_size * 0.5:
                end = start + cut + 1

        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)

        if end >= n:
            break
        start = max(end - overlap, start + 1)

    return chunks


class RAG:
    def __init__(self, lesson_path: str, client: OpenAI, cost_tracker=None):
        self.lesson_path = Path(lesson_path)
        self.db_path     = str(self.lesson_path / "chroma_db")
        self.client      = client

        self.ef = embedding_functions.OpenAIEmbeddingFunction(
            api_key=client.api_key,
            api_base="https://openrouter.ai/api/v1",
            model_name=MODEL_EMBEDDING
        )

        self.cost_tracker = cost_tracker
        self.chroma     = chromadb.PersistentClient(path=self.db_path)
        self.collection = self.chroma.get_or_create_collection(
            name="lesson",
            embedding_function=self.ef
        )

    def index_files(self):
        # PDF ถูกแปลงเป็น .txt ด้วย pdf_to_txt() ก่อนหน้านี้แล้ว จึงไม่รวม .pdf
        supported = {".txt", ".md", ".docx", ".pptx"}
        files = [
            f for f in self.lesson_path.iterdir()
            if f.is_file() and f.suffix.lower() in supported
        ]

        if not files:
            print("  ไม่พบไฟล์ที่รองรับในโฟลเดอร์นี้")
            return

        total_chunks = 0
        for file in files:
            print(f"  กำลังประมวลผล: {file.name}")
            try:
                text   = extract_text(str(file), self.client, self.cost_tracker)
                chunks = chunk_text(text)

                ids, docs, metas = [], [], []
                for i, chunk in enumerate(chunks):
                    chunk_id = hashlib.md5(f"{file.name}:{i}".encode()).hexdigest()
                    ids.append(chunk_id)
                    docs.append(chunk)
                    metas.append({"source": file.name, "chunk": i})

                batch = 50
                for b in range(0, len(ids), batch):
                    self.collection.upsert(
                        ids=ids[b:b+batch],
                        documents=docs[b:b+batch],
                        metadatas=metas[b:b+batch]
                    )

                total_chunks += len(chunks)
                print(f"    → {len(chunks)} chunks")

            except Exception as e:
                print(f"    ⚠️  ข้ามไฟล์ {file.name}: {e}")

        if total_chunks == 0:
            print("\n  ⚠️  ไม่มี chunk ถูกเก็บเลย — ตรวจสอบว่าไฟล์อ่านข้อความได้หรือไม่")
        else:
            print(f"\n  รวม {total_chunks} chunks เก็บใน ChromaDB แล้ว ✅")

        return total_chunks

    def query(self, question: str, n_results: int = 5) -> list[str]:
        count = self.collection.count()
        if count == 0:
            return []
        results = self.collection.query(
            query_texts=[question],
            n_results=min(n_results, count)
        )
        return results["documents"][0] if results["documents"] else []